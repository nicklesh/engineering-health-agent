"""Agent 10 - Executive Report.  Spec: agents/10_executive_report_agent.md

Builds, from validated outputs only:
  * output/presentation/engineering_health_review.pptx  (14 slides, native charts, speaker notes)
  * output/reports/executive_report.md
and then checks them (consistency_checks):
  * every number printed in the deck / report comes from the fact sheet (src/reporting/facts.py)
  * every chart series equals its source series
  * key facts recomputed independently from the raw CSVs, and compared with the dashboard data
    (raw <-> analysis <-> dashboard <-> deck)

Run:  python -m src.reporting.agent10_executive_report
"""
import json
import re
from statistics import mean, median

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from src.analysis.agent02_metrics_analysis import load_series
from src.analysis.dataset import load_clean_data, load_metric_catalog, pooled_ratio
from src.common.envelope import AGENTS, envelope_path, processed_dir, read_envelope, write_envelope
from src.config import ROOT, load_json, project_path
from src.data.loader import read_dataset
from src.reporting import agent09_dashboard
from src.reporting.agent09_dashboard import rolling
from src.reporting.facts import NUMBER, FactSheet, change, num, val

AGENT = AGENTS["10"]
# Stat cards show the statistic the trend test used. Saying so prevents an apparent conflict with
# tables of weekly averages (e.g. a 4.5 d median vs a 4.8 d mean for the same weeks).
STAT_NOTE = "Changes compare weeks 13–16 with weeks 1–6 (medians; pooled rates for percentages), as tested."
OUTPUT_DIR = project_path("output")  # tests point this at a temporary directory


def deck_path():
    return OUTPUT_DIR / "presentation" / "engineering_health_review.pptx"


def report_path():
    return OUTPUT_DIR / "reports" / "executive_report.md"

# ---- visual system: ink + white, deep teal accent, semantic red/green/amber (matches the dashboard)
INK, PAPER, PANEL, LINE, MUTED = "16191D", "FFFFFF", "F3F3F1", "DEDBD5", "666B73"
TEAL, TEAL_SOFT = "0E6B63", "DCECEA"
BAD, BAD_SOFT, GOOD, GOOD_SOFT, WARN, WARN_SOFT = "B4361F", "F6E2DC", "2E7D4F", "DFEFE4", "9A6A0C", "F4EAD2"
SEV = {"CRITICAL": "A3271A", "HIGH": "C4611B", "MEDIUM": "9A6A0C", "LOW": "6B7480"}
HEAD, BODY = "Cambria", "Calibri"
W, H = 13.333, 7.5
X0, CW = 0.6, 12.13  # left margin, content width


def rgb(h):
    return RGBColor.from_string(h)


# ======================================================================== facts
def build_facts():
    env = {k: read_envelope(AGENTS[k]) for k in ("01", "02", "03", "04", "05", "06", "07", "08")}
    series = load_series()
    store = series["scopes"]
    catalog = load_metric_catalog()
    thresholds = load_json(project_path("config/thresholds.json"))
    v08 = {f["id"]: f for f in env["08"]["findings"]}
    ok = {fid for fid, f in v08.items() if f["validation"]["verdict"] != "REJECT"}
    cls = {(r["scope"], r["metric"]): r for r in env["04"]["series_classifications"]}
    F = FactSheet()
    ctx = {"env": env, "store": store, "catalog": catalog, "v08": v08, "ok": ok, "cls": cls, "series": series, "th": thresholds}

    ws = series["week_starts"]
    F.add("period.first", ws[0], ws[0], "metric_series.week_starts[0]")
    F.add("period.last", ws[-1], ws[-1], "metric_series.week_starts[-1]")
    dq = env["01"]
    F.add("dq.score", dq["data_quality_score"], num(dq["data_quality_score"]), "01_data_validation.data_quality_score")
    F.add("dq.records", dq["validated_records"], f"{dq['validated_records']:,}", "01_data_validation.validated_records")
    F.add("dq.total", dq["summary"]["total_records"], f"{dq['summary']['total_records']:,}", "01_data_validation.summary.total_records")
    v = env["08"]["summary"]["verdicts"]
    for k in ("PASS", "WARN", "REJECT"):
        F.add(f"verdict.{k}", v[k], str(v[k]), f"08_evidence_validation.summary.verdicts.{k}")
    F.add("verdict.total", sum(v.values()), str(sum(v.values())), "08_evidence_validation.summary.validated")

    # KPIs (4-week rolling mean of Agent 02's weekly health index, as on the dashboard)
    health = env["02"]["health_scores"]["scopes"]["org"]
    kpi = {}
    for k in ("overall", "delivery", "quality", "reliability", "efficiency"):
        r = rolling([w[k] for w in health])
        kpi[k] = r
        F.add(f"kpi.{k}.now", r[-1], num(r[-1]), f"02 health_scores org {k}, 4-week rolling, week 16")
        F.add(f"kpi.{k}.delta", round(r[-1] - r[0], 1), f"{r[-1] - r[0]:+.1f}", f"02 health_scores org {k}, rolling W16 - W1")
    ctx["kpi"] = kpi

    def trend(fid, prefix):
        f = v08[fid]
        assert fid in ok, f"{fid} was rejected and must not be used"
        c = f["evidence"]["computed"]
        m = catalog[f["metrics"][0]]
        F.add(f"{prefix}.base", c["baseline"], val(c["baseline"], m["unit"]), f"{fid}.evidence.computed.baseline")
        F.add(f"{prefix}.recent", c["recent"], val(c["recent"], m["unit"]), f"{fid}.evidence.computed.recent")
        F.add(f"{prefix}.chg", c.get("change_pct"), change(c["baseline"], c["recent"], m["unit"], c.get("change_pct")),
              f"{fid}.evidence.computed.change(_pct)")
        F.add(f"{prefix}.since", f["weeks"][0], f"week {f['weeks'][0]}", f"{fid}.weeks[0]")
        F.add(f"{prefix}.conf", f["confidence"]["score"], f"{round(f['confidence']['score'] * 100)}%", f"{fid}.confidence.score (validated)")
        return f

    for metric, key in [("CYCLE_TIME_DAYS", "cycle"), ("PR_REVIEW_TIME_HOURS", "review"), ("DEFECTS", "defects"),
                        ("TEST_COVERAGE_PCT", "coverage"), ("LEAD_TIME_DAYS", "lead"), ("AUTOMATED_TEST_PCT", "auto")]:
        trend(f"TRD-ATLAS-ALL-{metric}", f"atlas.{key}")
    for metric, key in [("TEST_COVERAGE_PCT", "coverage"), ("CYCLE_TIME_DAYS", "cycle"), ("CHANGE_FAILURE_RATE_PCT", "cfr"),
                        ("DEPLOYMENT_SUCCESS_RATE_PCT", "success"), ("AUTOMATED_TEST_PCT", "auto")]:
        trend(f"TRD-PHOENIX-ALL-{metric}", f"phoenix.{key}")
    for metric, key in [("CHANGE_FAILURE_RATE_PCT", "cfr"), ("ROLLBACK_RATE_PCT", "rollback"), ("INCIDENTS", "incidents"),
                        ("FAILED_DEPLOYMENTS", "failed"), ("ON_CALL_PAGES", "oncall")]:
        trend(f"TRD-ALL-DATA-{metric}", f"data.{key}")
    trend("TRD-TITAN-ALL-BUILD_TIME_MIN", "titan.build")

    esc = cls[("team:Phoenix", "escaped_defect_rate_pct")]
    (x0, n0), (x1, n1) = esc["pooled_counts"]["baseline"], esc["pooled_counts"]["recent"]
    F.add("phoenix.esc.base", esc["baseline"], val(esc["baseline"], "%"), "04 classification team:Phoenix escaped baseline")
    F.add("phoenix.esc.recent", esc["recent"], val(esc["recent"], "%"), "04 classification team:Phoenix escaped recent")
    F.add("phoenix.esc.counts", [x0, n0, x1, n1], f"{x0}/{n0} vs {x1}/{n1}", "04 classification pooled_counts")
    F.add("phoenix.esc.p", esc["p_value"], num(esc["p_value"], 2), "04 classification team:Phoenix escaped p_value")

    h = v08["RISK-ALL-DATA-RELIABILITY-H-CHANGE_RELATED"]
    hc = h["evidence"]["computed"]
    F.add("data.chg_rel.base", hc["baseline"], f"{round(100 * hc['baseline'][0] / hc['baseline'][1])}%", f"{h['id']}.computed.baseline")
    F.add("data.chg_rel.recent", hc["recent"], f"{round(100 * hc['recent'][0] / hc['recent'][1])}%", f"{h['id']}.computed.recent")
    F.add("data.chg_rel.p", hc["p_value"], num(hc["p_value"], 3), f"{h['id']}.computed.p_value")
    F.add("data.chg_rel.verdict", h["validation"]["verdict"], h["validation"]["verdict"], f"{h['id']}.validation.verdict")

    for metric in ("CI_FAILURE_RATE_PCT", "ON_CALL_PAGES"):
        a = v08[f"ANO-NOVA-MOBILE-{metric}-W11"]
        m = catalog[a["metrics"][0]]
        F.add(f"nova.{metric.lower()}.value", a["evidence"]["computed"]["value"], val(a["evidence"]["computed"]["value"], m["unit"]), f"{a['id']}.computed.value")
        F.add(f"nova.{metric.lower()}.expected", a["expected_baseline"], val(a["expected_baseline"], m["unit"]), f"{a['id']}.expected_baseline")

    # Policy: a REJECTED finding never appears as a conclusion. One may appear only as a clearly
    # labelled rejection on the evidence-model slide, and its facts use the "rejected_example."
    # prefix so tests can enforce that (tests/test_reporting.py).
    rej = v08["RISK-TITAN-ALL-ENGINEERING_EFFICIENCY-H-THROUGHPUT_ITEMS"]
    assert rej["validation"]["verdict"] == "REJECT", "the rejected example must actually be rejected"
    d = rej["evidence"]["computed"]["common_drivers"][0]
    F.add("rejected_example.r", rej["evidence"]["computed"]["r"], num(rej["evidence"]["computed"]["r"], 2), f"{rej['id']}.computed.r")
    F.add("rejected_example.partial", d["partial_r"], f"{d['partial_r']:.2f}", f"{rej['id']}.computed.common_drivers[0].partial_r")
    F.add("rejected_example.hc_r", d["r_with_second"], num(d["r_with_second"], 2), f"{rej['id']}.computed.common_drivers[0].r_with_second")
    hcc = cls[("team:Titan", "team_headcount")]
    F.add("titan.hc.base", hcc["baseline"], num(hcc["baseline"]), "04 classification team:Titan team_headcount baseline")
    F.add("titan.hc.recent", hcc["recent"], num(hcc["recent"]), "04 classification team:Titan team_headcount recent")

    risks = sorted((f for f in env["08"]["findings"] if f["type"] == "risk" and f["id"] in ok), key=lambda f: f["rank"])
    for r in risks:
        F.add(f"risk.{r['id']}.conf", r["confidence"]["score"], f"{round(r['confidence']['score'] * 100)}%", f"{r['id']}.confidence.score (validated)")
    F.add("risk.count", len(risks), str(len(risks)), "08 risks not rejected")
    F.add("risk.critical", sum(r["severity"] == "CRITICAL" for r in risks), str(sum(r["severity"] == "CRITICAL" for r in risks)), "08 risks CRITICAL")
    recs = [f for f in env["08"]["findings"] if f["type"] == "recommendation" and f["id"] in ok]
    order = {"P1": 0, "P2": 1, "P3": 2}
    recs.sort(key=lambda r: (order[r["priority"]], -r["confidence"]["score"], r["id"]))
    F.add("rec.p1", sum(r["priority"] == "P1" for r in recs), str(sum(r["priority"] == "P1" for r in recs)), "08 recommendations P1")
    for r in recs:
        ms = r["measurement_of_success"]
        if ms["baseline"] is not None:
            m = catalog[ms["metric"]]
            F.add(f"rec.{r['id']}.base", ms["baseline"], val(ms["baseline"], m["unit"]), f"{r['id']}.measurement_of_success.baseline")
            F.add(f"rec.{r['id']}.now", ms["current"], val(ms["current"], m["unit"]), f"{r['id']}.measurement_of_success.current")
        F.add(f"rec.{r['id']}.target", ms["target"], ms["target"].split(" (")[0].replace("<=", "≤").replace(">=", "≥"),
              f"{r['id']}.measurement_of_success.target")
        F.add(f"rec.{r['id']}.by", ms["by_week"], f"week {ms['by_week']}", f"{r['id']}.measurement_of_success.by_week")
    ctx["risks"], ctx["recs"] = risks, recs

    # Team comparison: mean of the last 4 weeks (same as the dashboard's 4W view)
    comp_metrics = ["cycle_time_days", "pr_review_time_hours", "test_coverage_pct", "defects", "change_failure_rate_pct", "incidents"]
    comp = {}
    for team in ("Atlas", "Nova", "Orion", "Phoenix", "Titan"):
        for mid in comp_metrics:
            vals = [x for x in store[f"team:{team}"]["metrics"][mid]["values"][-4:] if x is not None]
            m = mean(vals)
            comp[(team, mid)] = m
            F.add(f"cmp.{team}.{mid}", round(m, 2), val(m, catalog[mid]["unit"]), f"metric_series team:{team} {mid} mean W13-16")
    ctx["comp"], ctx["comp_metrics"] = comp, comp_metrics

    t = thresholds["trend"]
    F.add("th.min_weeks", t["min_trend_weeks"], str(t["min_trend_weeks"]), "thresholds.trend.min_trend_weeks")
    F.add("th.min_pct", t["min_change_pct"], num(t["min_change_pct"]), "thresholds.trend.min_change_pct")
    F.add("th.min_pp", t["min_change_pp"], num(t["min_change_pp"]), "thresholds.trend.min_change_pp")
    F.add("th.effect", t["min_effect_sd"], num(t["min_effect_sd"]), "thresholds.trend.min_effect_sd")
    F.add("th.p", t["max_p_value"], num(t["max_p_value"], 2), "thresholds.trend.max_p_value")
    F.add("th.tau", t["min_abs_tau"], num(t["min_abs_tau"]), "thresholds.trend.min_abs_tau")
    F.add("th.partial", thresholds["correlation"]["partial_r_explained_below"], num(thresholds["correlation"]["partial_r_explained_below"]), "thresholds.correlation.partial_r_explained_below")
    F.add("th.anomaly_z", thresholds["anomaly"]["min_z"], num(thresholds["anomaly"]["min_z"]), "thresholds.anomaly.min_z")
    for k, w in thresholds["confidence"]["weights"].items():
        F.add(f"conf.w.{k}", w, f"{round(w * 100)}%", f"thresholds.confidence.weights.{k}")
    F.add("health.base", 75, "75", "thresholds.metrics_analysis.health_score_at_baseline")
    F.add("health.better", 100, "100", "health index: 25%+ better than baseline")
    F.add("health.worse", 50, "50", "health index: 25% worse than baseline")
    F.add("health.band", 25, "25%", "health index scale")
    return F, ctx


# ======================================================================== drawing helpers
class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.blank = self.prs.slide_layouts[6]
        self.chart_sources = []
        self.n = 0

    def slide(self, dark=False, notes=""):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(INK if dark else PAPER)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        if not dark:
            self.text(s, X0, H - 0.42, 9, 0.25, "Engineering Health Review · synthetic organisation · validated analysis", 10, MUTED)
            self.text(s, W - X0 - 1, H - 0.42, 1, 0.25, str(self.n), 10, MUTED, align=PP_ALIGN.RIGHT)
        return s

    def text(self, s, x, y, w, h, content, size=14, color=INK, bold=False, font=BODY, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, italic=False, spacing=None):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        paras = content if isinstance(content, list) else [content]
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            if spacing:
                para.space_after = Pt(spacing)
            runs = p if isinstance(p, list) else [(p, {})]
            for rt in runs:
                txt, o = (rt, {}) if isinstance(rt, str) else rt
                r = para.add_run()
                r.text = txt
                f = r.font
                f.name = o.get("font", font)
                f.size = Pt(o.get("size", size))
                f.bold = o.get("bold", bold)
                f.italic = o.get("italic", italic)
                f.color.rgb = rgb(o.get("color", color))
        return tb

    def rect(self, s, x, y, w, h, fill, line=None, rounded=False):
        shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
        if rounded:
            shp.adjustments[0] = 0.08
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        if line:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(0.75)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def dot(self, s, x, y, d, fill):
        shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def title(self, s, eyebrow, title, dark=False):
        self.text(s, X0, 0.42, CW, 0.3, eyebrow.upper(), 11, TEAL if not dark else "7FC8BE", bold=True)
        self.text(s, X0, 0.72, CW, 0.95, title, 30, PAPER if dark else INK, bold=True, font=HEAD)

    def stat(self, s, x, y, w, label, value, sub, color=INK, sub_color=MUTED, panel=PANEL):
        self.rect(s, x, y, w, 1.35, panel, rounded=True)
        self.text(s, x + 0.2, y + 0.14, w - 0.4, 0.25, label, 11, MUTED, bold=True)
        self.text(s, x + 0.2, y + 0.38, w - 0.4, 0.55, value, 26, color, bold=True, font=HEAD)
        self.text(s, x + 0.2, y + 0.94, w - 0.4, 0.3, sub, 12, sub_color, bold=True)

    @staticmethod
    def axis_range(series):
        """Fit the value axis to the data (with padding) so trends are not flattened against zero."""
        vals = [v for _, values, _, _ in series for v in values if v is not None]
        lo, hi = min(vals), max(vals)
        span = (hi - lo) or abs(hi) or 1
        step = next(s for s in (0.5, 1, 2, 5, 10, 20, 25, 50) if span / s <= 6)
        lo_n = max(0.0, (int((lo - span * 0.1) / step)) * step) if lo >= 0 else (int((lo - span * 0.1) / step) - 1) * step
        hi_n = (int((hi + span * 0.1) / step) + 1) * step
        return lo_n, hi_n, step

    def line_chart(self, s, x, y, w, h, cats, series, number_format="0", source=""):
        cd = CategoryChartData()
        cd.categories = cats
        for name, values, _, _ in series:
            cd.add_series(name, values)
        gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        ch = gf.chart
        ch.has_title = False
        ch.font.name = BODY
        ch.font.size = Pt(11)
        ch.font.color.rgb = rgb(MUTED)
        ch.has_legend = len(series) > 1
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.TOP
            ch.legend.include_in_layout = False
            ch.legend.font.size = Pt(11)
        va = ch.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = rgb(LINE)
        va.format.line.fill.background()
        va.tick_labels.number_format = number_format
        va.tick_labels.number_format_is_linked = False
        va.minimum_scale, va.maximum_scale, va.major_unit = self.axis_range(series)
        ca = ch.category_axis
        ca.format.line.color.rgb = rgb(LINE)
        ca.tick_labels.font.size = Pt(10)
        for i, (name, values, color, width) in enumerate(series):
            ser = ch.plots[0].series[i]
            ser.smooth = False
            ser.format.line.color.rgb = rgb(color)
            ser.format.line.width = Pt(width)
            ser.marker.style = XL_MARKER_STYLE.CIRCLE
            ser.marker.size = 4 if width >= 2 else 3
            ser.marker.format.fill.solid()
            ser.marker.format.fill.fore_color.rgb = rgb(color)
            ser.marker.format.line.color.rgb = rgb(color)
            self.chart_sources.append({"slide": self.n, "chart": gf.name, "series": name, "values": list(values), "source": source})
        return gf

    def bar_chart(self, s, x, y, w, h, cats, values, color, number_format, source):
        cd = CategoryChartData()
        cd.categories = cats
        cd.add_series("value", values)
        gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        ch = gf.chart
        ch.has_title = ch.has_legend = False
        ch.font.name = BODY
        ch.font.size = Pt(11)
        ch.font.color.rgb = rgb(INK)
        plot = ch.plots[0]
        plot.gap_width = 60
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = number_format
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size = Pt(11)
        ser = plot.series[0]
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(color)
        ch.value_axis.visible = False
        ch.value_axis.has_major_gridlines = False
        ch.category_axis.format.line.fill.background()
        ch.category_axis.reverse_order = True
        self.chart_sources.append({"slide": self.n, "chart": gf.name, "series": "value", "values": list(values), "source": source})
        return gf

    def table(self, s, x, y, w, col_w, rows, header_fill=INK, row_h=0.42, font_size=12, fills=None):
        shp = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
        tbl = shp.table
        tbl.first_row = True
        tbl.horz_banding = False
        for j, cw in enumerate(col_w):
            tbl.columns[j].width = Inches(cw)
        for i, row in enumerate(rows):
            tbl.rows[i].height = Inches(row_h)
            for j, value in enumerate(row):
                cell = tbl.cell(i, j)
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                fill = header_fill if i == 0 else (fills or {}).get((i, j), PAPER if i % 2 else PANEL)
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(fill)
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                txt, opts = (value, {}) if isinstance(value, str) else value
                r = p.add_run()
                r.text = txt
                r.font.name = BODY
                r.font.size = Pt(opts.get("size", font_size if i else font_size - 1))
                r.font.bold = opts.get("bold", i == 0)
                r.font.color.rgb = rgb(opts.get("color", PAPER if i == 0 else INK))
                p.alignment = opts.get("align", PP_ALIGN.LEFT)
        return shp

    def arrow(self, s, x1, y1, x2, y2, color=MUTED):
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = rgb(color)
        c.line.width = Pt(1.5)
        ln = c.line._get_or_add_ln()
        tail = ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)
        return c


# ======================================================================== slides
def build_deck(F, ctx):
    D = Deck()
    store, catalog, v08, cls = ctx["store"], ctx["catalog"], ctx["v08"], ctx["cls"]
    weeks = [f"W{i}" for i in range(1, 17)]
    risks, recs = ctx["risks"], ctx["recs"]
    p1 = [r for r in recs if r["priority"] == "P1"]

    # 1 ---- title (dark)
    s = D.slide(dark=True, notes=(
        "Purpose: decide on two critical risks and endorse four actions. Everything in this deck comes from an 8-agent analysis "
        "pipeline over 16 weeks of synthetic engineering data. Every number was validated against its source records before "
        "it reached a slide."))
    D.text(s, X0, 0.9, 8, 0.3, "ENGINEERING HEALTH REVIEW", 12, "7FC8BE", bold=True)
    D.text(s, X0, 1.35, 7.6, 2.4, "Two risks need leadership attention, and one team shows the way forward", 40, PAPER, bold=True, font=HEAD)
    D.text(s, X0, 4.15, 7.4, 0.9, [f"Weeks 1–16 · {F['period.first']} to {F['period.last']} (week starts)",
                                   "Synthetic organisation · 5 teams · 4 platforms"], 16, "C9CDD2", spacing=4)
    D.rect(s, 8.85, 1.35, 3.85, 3.2, "23272C", rounded=True)
    D.text(s, 9.15, 1.6, 3.3, 0.3, "ENGINEERING HEALTH", 11, "9EA4AB", bold=True)
    D.text(s, 9.15, 1.95, 3.3, 1.0, [[(F["kpi.overall.now"], {"size": 60, "bold": True, "font": HEAD, "color": PAPER}),
                                      (" / 100", {"size": 18, "color": "9EA4AB"})]], 60)
    D.text(s, 9.15, 3.15, 3.3, 0.35, f"{F['kpi.overall.delta']} points since week 1", 16, "F0A08F", bold=True)
    D.text(s, 9.15, 3.6, 3.3, 0.7, f"Relative index: {F['health.base']} = organisation level in weeks 1–4", 12, "9EA4AB")
    D.text(s, X0, 6.55, CW, 0.3, f"Data quality {F['dq.score']}/100 · {F['verdict.total']} findings checked: {F['verdict.PASS']} passed, "
                                 f"{F['verdict.WARN']} with caveats, {F['verdict.REJECT']} rejected", 12, "9EA4AB")

    # 2 ---- executive summary
    s = D.slide(notes=(
        "Four messages. Two critical risks: Atlas is slowing down and losing quality at the same time; the Data platform's "
        "releases fail far more often. Phoenix improved on every quality and delivery measure we track, so its practices are worth spreading. "
        "The Nova week-11 spike was a one-off that recovered the next week; it needs a post-incident check, not a programme."))
    D.title(s, "Executive summary", f"Health is down {F['kpi.overall.delta'].lstrip('+-')} points, driven by two critical risks")
    rows = [
        (SEV["CRITICAL"], "Atlas is slowing down and losing quality",
         f"Cycle time {F['atlas.cycle.base']} → {F['atlas.cycle.recent']} ({F['atlas.cycle.chg']}) · defects {F['atlas.defects.base']} → {F['atlas.defects.recent']} per week · coverage {F['atlas.coverage.base']} → {F['atlas.coverage.recent']}"),
        (SEV["CRITICAL"], "Data platform releases fail far more often",
         f"Change failure rate {F['data.cfr.base']} → {F['data.cfr.recent']} · rollbacks {F['data.rollback.base']} → {F['data.rollback.recent']} · incidents {F['data.incidents.base']} → {F['data.incidents.recent']} per week"),
        (GOOD, "Phoenix improved across quality and delivery",
         f"Coverage {F['phoenix.coverage.base']} → {F['phoenix.coverage.recent']} · cycle time {F['phoenix.cycle.base']} → {F['phoenix.cycle.recent']} · change failure rate {F['phoenix.cfr.base']} → {F['phoenix.cfr.recent']}"),
        (SEV["LOW"], "Nova's week-11 spike was a one-off",
         f"CI failure rate {F['nova.ci_failure_rate_pct.value']} against an expected {F['nova.ci_failure_rate_pct.expected']}; back to normal the next week. No structural action."),
    ]
    y = 1.95
    for color, head, detail in rows:
        D.dot(s, X0, y + 0.06, 0.26, color)
        D.text(s, X0 + 0.45, y, 7.3, 0.35, head, 18, INK, bold=True)
        D.text(s, X0 + 0.45, y + 0.4, 7.3, 0.6, detail, 14, MUTED)
        y += 1.18
    D.text(s, X0, 6.72, 7.8, 0.3, STAT_NOTE, 10, MUTED)
    D.rect(s, 8.75, 1.95, 3.98, 4.75, PANEL, rounded=True)
    D.text(s, 9.0, 2.15, 3.5, 0.3, "WHAT WE ASK", 11, TEAL, bold=True)
    D.text(s, 9.0, 2.5, 3.5, 0.45, f"Endorse {F['rec.p1']} priority actions", 18, INK, bold=True, font=HEAD)
    D.text(s, 9.0, 3.1, 3.5, 3.4, [[("P1  ", {"bold": True, "color": TEAL}), (r["headline"], {})] for r in p1], 14, INK, spacing=10)

    # 3 ---- health score
    s = D.slide(notes=(
        "The health index is relative: 75 means the organisation's own level in weeks 1 to 4; 100 is 25% or more better, 50 is 25% worse. "
        "Cards show a 4-week rolling average. Delivery fell most, mainly from Atlas's slower flow and the Data platform's failed releases."))
    deltas = {k: float(F.facts[f"kpi.{k}.delta"].value) for k in ("delivery", "quality", "reliability", "efficiency")}
    worst = min(deltas, key=deltas.get)
    D.title(s, "Engineering health score", f"Health is down {F['kpi.overall.delta'].lstrip('+-')} points since week 1; {worst} fell most")
    labels = {"overall": "Engineering health", "delivery": "Delivery", "quality": "Quality", "reliability": "Reliability", "efficiency": "Efficiency"}
    x, cw = X0, (CW - 4 * 0.25) / 5
    for k in labels:
        dv = float(F.facts[f"kpi.{k}.delta"].value)
        D.stat(s, x, 1.85, cw, labels[k].upper(), F[f"kpi.{k}.now"], f"{F[f'kpi.{k}.delta']} since W1",
               color=INK, sub_color=BAD if dv < 0 else GOOD, panel=PANEL)
        x += cw + 0.25
    kp = ctx["kpi"]
    D.line_chart(s, X0, 3.45, CW, 3.45, weeks,
                 [("Engineering health", kp["overall"], INK, 2.75), ("Delivery", kp["delivery"], BAD, 1.5),
                  ("Quality", kp["quality"], WARN, 1.5), ("Reliability", kp["reliability"], TEAL, 1.5),
                  ("Efficiency", kp["efficiency"], "8A9098", 1.5)], "0", "02 health_scores org, 4-week rolling")

    # 4 ---- what improved
    s = D.slide(notes=(
        f"Phoenix improved on coverage, automation, cycle time and release reliability. Its escaped-defect rate also fell, from "
        f"{F['phoenix.esc.base']} to {F['phoenix.esc.recent']}, but that rests on {F['phoenix.esc.counts']} defects (p = {F['phoenix.esc.p']}). "
        "The system does not call that a trend yet, and neither should we."))
    D.title(s, "What improved", "Phoenix shows what good looks like")
    stats = [("TEST COVERAGE", "phoenix.coverage"), ("CYCLE TIME", "phoenix.cycle"),
             ("CHANGE FAILURE RATE", "phoenix.cfr"), ("DEPLOYMENT SUCCESS", "phoenix.success")]
    for i, (label, key) in enumerate(stats):
        cx, cy = X0 + (i % 2) * 2.95, 1.9 + (i // 2) * 1.6
        D.stat(s, cx, cy, 2.75, label, F[f"{key}.recent"], f"from {F[f'{key}.base']} ({F[f'{key}.chg']})", sub_color=GOOD)
    D.text(s, X0, 4.98, 5.6, 0.4, STAT_NOTE, 10, MUTED)
    D.text(s, X0, 5.55, 5.6, 1.0, [[("Directional only: ", {"bold": True, "color": WARN}),
                                   (f"escaped-defect rate {F['phoenix.esc.base']} → {F['phoenix.esc.recent']} ({F['phoenix.esc.counts']} defects, "
                                    f"p = {F['phoenix.esc.p']}). Too few defects to call it a trend yet.", {})]], 13, INK)
    cov = store["team:Phoenix"]["metrics"]["test_coverage_pct"]["values"]
    org = store["org"]["metrics"]["test_coverage_pct"]["values"]
    D.text(s, 6.85, 1.9, 5.9, 0.3, "Test coverage, % (weekly)", 12, MUTED, bold=True)
    D.line_chart(s, 6.75, 2.15, 5.98, 4.45, weeks, [("Phoenix", cov, GOOD, 2.75), ("Organisation", org, "9AA0A7", 1.5)],
                 "0", "metric_series team:Phoenix / org test_coverage_pct")

    # 5 ---- what deteriorated
    s = D.slide(notes=(
        "At Atlas, flow and quality degrade together. Review time and cycle time rose while coverage and automation fell and defects doubled. "
        "This is association, not a diagnosis. The leading-indicator analysis says coverage changes preceded defect changes by one week. "
        "That makes coverage the first place to look, not a proven cause."))
    D.title(s, "What deteriorated", f"Atlas: slower delivery and falling quality, both since {F['atlas.review.since']}")
    items = [("CYCLE TIME", "atlas.cycle"), ("PR REVIEW TIME", "atlas.review"), ("DEFECTS PER WEEK", "atlas.defects"), ("TEST COVERAGE", "atlas.coverage")]
    for i, (label, key) in enumerate(items):
        cx, cy = X0 + (i % 2) * 2.95, 1.9 + (i // 2) * 1.6
        D.stat(s, cx, cy, 2.75, label, F[f"{key}.recent"], f"from {F[f'{key}.base']} ({F[f'{key}.chg']})", sub_color=BAD)
    D.text(s, X0, 4.98, 5.6, 0.4, STAT_NOTE, 10, MUTED)
    D.text(s, X0, 5.55, 5.6, 1.0, [[("Read with care: ", {"bold": True, "color": TEAL}),
                                   ("these move together, which is association. Coverage changes preceded defect changes by one week, so "
                                    "coverage is the first place to look, not a proven cause.", {})]], 13, INK)
    D.text(s, 6.85, 1.9, 5.9, 0.3, "Cycle time, days (weekly)", 12, MUTED, bold=True)
    D.line_chart(s, 6.75, 2.15, 5.98, 4.45, weeks,
                 [("Atlas", store["team:Atlas"]["metrics"]["cycle_time_days"]["values"], BAD, 2.75),
                  ("Organisation", store["org"]["metrics"]["cycle_time_days"]["values"], "9AA0A7", 1.5)],
                 "0.0", "metric_series team:Atlas / org cycle_time_days")

    # 6 ---- emerging trends: onset timeline
    s = D.slide(notes=(
        f"Each bar starts in the week a validated trend began and runs to week 16. A trend needs at least {F['th.min_weeks']} weeks, a significant "
        f"monotonic change, at least {F['th.min_pct']}% (or {F['th.min_pp']} percentage points for rates), and an effect of at least {F['th.effect']} times "
        "normal noise. That is why the week-11 spike at Nova does not appear here."))
    # The trends behind each validated risk (strongest first), plus the trends behind Phoenix's
    # improvement - so the timeline shows exactly what the deck's conclusions rest on.
    picks, seen = [], set()
    groups = [([r["id"] for r in risks if r["entity"] == "team:Atlas"], 4), ([r["id"] for r in risks if r["entity"] == "platform:Data"], 4),
              ([r["id"] for r in risks if r["entity"] == "team:Titan"], 1), (["QUA-PHOENIX-ALL-STATUS"], 3)]
    for ids, n in groups:
        refs = [ref for rid in ids for ref in (v08[rid].get("primary_refs") or v08[rid]["evidence"].get("finding_refs", []))]
        fs = [v08[ref] for ref in refs if ref in ctx["ok"] and v08[ref]["type"] == "trend" and ref not in seen]
        fs.sort(key=lambda f: (-f["confidence"]["score"], f["id"]))
        picks += fs[:n]
        seen.update(f["id"] for f in fs[:n])
    # Title claim computed from the trends shown (and re-verified in consistency_checks).
    bad_onsets = sorted(f["weeks"][0] for f in picks if f["classification"] == "DETERIORATING")
    F.add("tl.first", bad_onsets[0], f"week {bad_onsets[0]}", "earliest onset among deteriorating trends shown")
    F.add("tl.last", bad_onsets[-1], f"week {bad_onsets[-1]}", "latest onset among deteriorating trends shown")
    ctx["timeline"] = picks
    D.title(s, "Emerging trends", f"Problems began between {F['tl.first']} and {F['tl.last']}; none is a blip")
    ax0, ax1, top = 4.1, 11.3, 2.15
    step = (ax1 - ax0) / 15
    rowh = min(0.4, 4.3 / len(picks))
    for w in range(1, 17, 3):
        D.text(s, ax0 + (w - 1) * step - 0.3, top - 0.32, 0.6, 0.25, f"W{w}", 10, MUTED, align=PP_ALIGN.CENTER)
    D.text(s, ax0 + 15 * step - 0.3, top - 0.32, 0.6, 0.25, "W16", 10, MUTED, align=PP_ALIGN.CENTER)
    for i, f in enumerate(picks):
        yy = top + i * rowh
        m = catalog[f["metrics"][0]]
        c = f["evidence"]["computed"]
        key = f"tl.{f['id']}"
        F.add(key + ".chg", c.get("change_pct"), change(c["baseline"], c["recent"], m["unit"], c.get("change_pct")), f"{f['id']}.evidence.computed")
        D.text(s, X0, yy + 0.04, 3.4, rowh - 0.06, f"{store[f['dimension']['scope']]['label']} · {m['label']}", 13, INK)
        onset = f["weeks"][0]
        color = GOOD if f["classification"] == "IMPROVING" else BAD
        D.rect(s, ax0 + (onset - 1) * step, yy + 0.08, (16 - onset) * step + 0.06, rowh - 0.18, color, rounded=True)
        D.text(s, ax1 + 0.2, yy + 0.04, 1.3, rowh - 0.06, F[key + ".chg"], 13, color, bold=True)
    D.text(s, X0, top + len(picks) * rowh + 0.2, CW, 0.3, "Bars run from the week each trend began to week 16. Red = deteriorating, green = improving.", 11, MUTED)

    # 7 ---- top risks table
    s = D.slide(notes="Risks are ranked by severity points (number of deteriorating metrics, size of change, tier-1 services, reliability impact), then by validated confidence. "
                      "A one-week disruption is capped at LOW by rule. Every risk passed evidence validation.")
    D.title(s, "Top engineering risks", f"{F['risk.count']} validated risks; {F['risk.critical']} are critical")
    rows = [["#", "Severity", "Risk", "Category", "Confidence", "Validation"]]
    fills = {}
    for i, r in enumerate(risks, start=1):
        rows.append([str(r["rank"]), (r["severity"], {"bold": True, "color": PAPER}), r["risk"].split(" (")[0], r["category"],
                     F[f"risk.{r['id']}.conf"], "Passed" if r["validation"]["verdict"] == "PASS" else "Caveat"])
        fills[(i, 1)] = SEV[r["severity"]]
    D.table(s, X0, 1.9, CW, [0.45, 1.35, 6.2, 1.9, 1.15, 1.08], rows, row_h=0.5, font_size=13, fills=fills)

    # 8 ---- quality & reliability (Data platform)
    s = D.slide(notes=(
        f"The Data platform's change failure rate rose from {F['data.cfr.base']} to {F['data.cfr.recent']}, with rollbacks and incidents following. "
        f"Post-incident reviews attributed more incidents to changes ({F['data.chg_rel.base']} to {F['data.chg_rel.recent']}), but with p = "
        f"{F['data.chg_rel.p']} this is a lead to check, not a conclusion. A tempting explanation, that rising on-call load drives rollbacks, "
        "was rejected: the two only share a time trend."))
    D.title(s, "Quality & reliability", f"Data platform: change failure rate rose from {F['data.cfr.base']} to {F['data.cfr.recent']}")
    entry_d = store["platform:Data"]["metrics"]["change_failure_rate_pct"]
    entry_o = store["org"]["metrics"]["change_failure_rate_pct"]
    pd = [pooled_ratio(entry_d, w, 4) for w in range(1, 17)]
    po = [pooled_ratio(entry_o, w, 4) for w in range(1, 17)]
    D.text(s, X0, 1.9, 7.2, 0.3, "Change failure rate, % (4-week pooled)", 12, MUTED, bold=True)
    D.line_chart(s, X0 - 0.1, 2.15, 7.3, 4.5, weeks, [("Data platform", pd, BAD, 2.75), ("Organisation", po, "9AA0A7", 1.5)],
                 "0", "metric_series platform:Data / org change_failure_rate_pct, 4-week pooled")
    for i, (label, key) in enumerate([("ROLLBACK RATE", "data.rollback"), ("INCIDENTS PER WEEK", "data.incidents"), ("ON-CALL PAGES PER WEEK", "data.oncall")]):
        D.stat(s, 8.1, 1.9 + i * 1.5, 4.63, label, F[f"{key}.recent"], f"from {F[f'{key}.base']} ({F[f'{key}.chg']})", sub_color=BAD)
    D.text(s, X0, 6.72, 7.2, 0.3, STAT_NOTE, 10, MUTED)
    D.text(s, 8.1, 6.3, 4.63, 0.6, [[("Possible explanation, unconfirmed: ", {"bold": True, "color": WARN}),
                                     (f"change-related incidents {F['data.chg_rel.base']} → {F['data.chg_rel.recent']} (p = {F['data.chg_rel.p']}).", {})]], 12, INK)

    # 9 ---- team comparison
    s = D.slide(notes="Values are averages over weeks 13 to 16. Red cells are at least 15% worse than the median team, green at least 15% better. "
                      "Atlas trails on flow and quality; Titan and Orion carry the Data platform's release failures; Phoenix leads.")
    D.title(s, "Team comparison", "Atlas trails on flow and quality; Data-platform teams on releases")
    cm = ctx["comp_metrics"]
    head = ["Team"] + [catalog[m]["label"] + ("" if catalog[m]["unit"] in ("count",) else f" ({ {'days': 'd', 'hours': 'h', '%': '%'}.get(catalog[m]['unit'], catalog[m]['unit']) })") for m in cm]
    rows, fills = [head], {}
    teams = ["Atlas", "Nova", "Orion", "Phoenix", "Titan"]
    for i, t in enumerate(teams, start=1):
        row = [(t, {"bold": True})]
        for j, m in enumerate(cm, start=1):
            vals = [ctx["comp"][(x, m)] for x in teams]
            med = median(vals)
            v = ctx["comp"][(t, m)]
            rel = ((med - v) if catalog[m]["polarity"] == "lower_is_better" else (v - med)) / abs(med) if med else 0
            fills[(i, j)] = BAD_SOFT if rel <= -0.15 else GOOD_SOFT if rel >= 0.15 else (PAPER if i % 2 else PANEL)
            row.append((F[f"cmp.{t}.{m}"], {"align": PP_ALIGN.RIGHT}))
        rows.append(row)
    D.table(s, X0, 1.95, CW, [1.6] + [(CW - 1.6) / len(cm)] * len(cm), rows, row_h=0.62, font_size=15, fills=fills)
    D.text(s, X0, 5.95, CW, 0.5, "Average of weeks 13–16. Red: at least 15% worse than the median team; green: at least 15% better. Defects and incidents are counts per week.", 12, MUTED)

    # 10 ---- recommended actions
    s = D.slide(notes="Each action names an owner and a measurable target from the evidence. P1 actions address the two critical risks and the high-confidence Atlas risks. "
                      "The Nova action is deliberately light: confirm the post-incident review and act only if the spike recurs.")
    D.title(s, "Recommended actions", f"{F['rec.p1']} priority actions, each with an owner and a measurable target")
    for i, r in enumerate(p1):
        cx, cy = X0 + (i % 2) * 6.17, 1.9 + (i // 2) * 2.05
        D.rect(s, cx, cy, 5.96, 1.85, PANEL, rounded=True)
        D.text(s, cx + 0.25, cy + 0.2, 5.4, 0.3, f"P1 · {r['owner_type'].upper()}", 11, TEAL, bold=True)
        D.text(s, cx + 0.25, cy + 0.5, 5.4, 0.75, r["headline"], 17, INK, bold=True, font=HEAD)
        ms = r["measurement_of_success"]
        D.text(s, cx + 0.25, cy + 1.3, 5.4, 0.4, f"Target: {catalog[ms['metric']]['label'].lower()} {F[f'rec.{r['id']}.target']} by {F[f'rec.{r['id']}.by']}", 13, MUTED)
    others = [r for r in recs if r["priority"] != "P1"]
    D.text(s, X0, 6.05, CW, 0.6, [[(f"{r['priority']}  ", {"bold": True, "color": TEAL}), (f"{r['headline']} ({r['owner_type']})", {})] for r in others], 12, INK, spacing=2)

    # 11 ---- expected outcomes / measurement
    s = D.slide(notes="Targets bring each lead metric back within one trend threshold of its own baseline. They are review points, not stretch goals. "
                      "If a target is missed at week 22, the risk stays open and the evidence is re-run.")
    D.title(s, "Expected outcomes & measurement", "Success means each lead metric is back near its own baseline by week 22")
    rows = [["Action", "Lead metric", "Baseline", "Now", "Target", "By", "Owner"]]
    for r in recs:
        ms = r["measurement_of_success"]
        if ms["baseline"] is None:
            continue
        rows.append([r["headline"], catalog[ms["metric"]]["label"], F[f"rec.{r['id']}.base"], F[f"rec.{r['id']}.now"],
                     F[f"rec.{r['id']}.target"], F[f"rec.{r['id']}.by"], r["owner_type"]])
    D.table(s, X0, 1.9, CW, [3.9, 1.75, 1.05, 1.05, 1.2, 0.95, 2.23], rows, row_h=0.6, font_size=12)

    # 12 ---- how the system works
    s = D.slide(notes="Ten narrow agents with explicit contracts. Code does all arithmetic; AI is used only for wording, and its wording is checked again. "
                      "The validation gate stops the run on bad data. The evidence validator can reject any finding, and recommendations built on a rejected finding are regenerated.")
    D.title(s, "How the intelligence system works", "Narrow agents, and a critic that can say no")
    boxes = [("Synthetic data", "5 CSV files", 0.6, 2.2), ("01 Validation gate", "stop on bad data", 2.75, 2.2),
             ("02 Metrics", "traceable series", 4.9, 2.2), ("04 Trends", "", 7.05, 1.75), ("05 Anomalies", "", 7.05, 2.65),
             ("03 Quality", "", 9.2, 2.2), ("06 Risk", "proposes widely", 9.2, 3.85), ("07 Coach", "actions + targets", 7.05, 3.85),
             ("08 Evidence validation", "pass / caveat / reject", 4.9, 3.85), ("09 Dashboard", "", 2.75, 5.35), ("10 Report + deck", "", 4.9, 5.35),
             ("AI narratives", "re-validated", 7.05, 5.35)]
    pos = {}
    for name, sub, bx, by in boxes:
        key = name.split(" ")[0]
        dark = name.startswith(("01", "08"))
        D.rect(s, bx, by, 1.95, 0.75, INK if dark else PANEL, rounded=True)
        D.text(s, bx + 0.12, by + 0.1, 1.71, 0.3, name, 12, PAPER if dark else INK, bold=True)
        if sub:
            D.text(s, bx + 0.12, by + 0.4, 1.71, 0.28, sub, 10, "C9CDD2" if dark else MUTED)
        pos[name] = (bx, by)
    for a, b in [((2.55, 2.575), (2.75, 2.575)), ((4.7, 2.575), (4.9, 2.575)), ((6.85, 2.45), (7.05, 2.125)), ((6.85, 2.7), (7.05, 3.025)),
                 ((9.0, 2.125), (9.2, 2.45)), ((9.0, 3.025), (9.2, 2.7)), ((10.175, 2.95), (10.175, 3.85)), ((9.2, 4.225), (9.0, 4.225)),
                 ((7.05, 4.225), (6.85, 4.225)), ((5.3, 4.6), (3.9, 5.35)), ((5.875, 4.6), (5.875, 5.35)), ((6.4, 4.6), (7.6, 5.35))]:
        D.arrow(s, a[0], a[1], b[0], b[1])
    D.text(s, X0, 3.95, 4.0, 1.2, [[("Feedback loop: ", {"bold": True, "color": TEAL}),
                                   ("a rejected risk sends Agent 07 back to rebuild its recommendations without it, then Agent 08 checks again.", {})]], 12, INK)
    D.text(s, 9.2, 5.3, 3.5, 1.3, [[("Rule: ", {"bold": True, "color": TEAL}),
                                   ("code computes every number; AI may only reword validated findings.", {})]], 12, INK)

    # 13 ---- evidence & confidence
    s = D.slide(notes=(
        f"The evidence validator re-computed every finding and rejected {F['verdict.REJECT']}. The most instructive rejection is shown: throughput and build time "
        f"at Titan move together (r = {F['rejected_example.r']}), but both follow headcount (r = {F['rejected_example.hc_r']}); after controlling for it the association is "
        f"{F['rejected_example.partial']}. Correlation is not causation. Confidence is computed from six named components, never typed in."))
    D.title(s, "Evidence & confidence model", "Every finding is challenged before it reaches this deck")
    for i, (label, key, color) in enumerate([("PASSED", "verdict.PASS", GOOD), ("WITH CAVEAT", "verdict.WARN", WARN), ("REJECTED", "verdict.REJECT", BAD)]):
        D.stat(s, X0 + i * 2.0, 1.9, 1.85, label, F[key], "findings", color=color)
    D.rect(s, X0, 3.55, 5.85, 2.55, BAD_SOFT, rounded=True)
    D.text(s, X0 + 0.25, 3.75, 5.35, 0.3, "REJECTED EXAMPLE", 11, BAD, bold=True)
    D.text(s, X0 + 0.25, 4.05, 5.35, 0.7, "“Rising throughput may be slowing Titan's builds”", 17, INK, bold=True, font=HEAD)
    D.text(s, X0 + 0.25, 4.85, 5.35, 1.5, f"Both track team size (r = {F['rejected_example.hc_r']}; headcount {F['titan.hc.base']} → {F['titan.hc.recent']}). "
                                          f"Controlling for headcount, the association drops to r = {F['rejected_example.partial']}. A common driver, not a cause.", 13, INK)
    wk = load_json(project_path("config/thresholds.json"))["confidence"]["weights"]
    names = {"completeness": "Data completeness", "observations": "Observations", "consistency": "Consistency", "magnitude": "Magnitude",
             "evidence_quality": "Evidence quality", "validation": "Validation verdict"}
    D.text(s, 6.95, 1.9, 5.8, 0.3, "How confidence is computed (component weights)", 12, MUTED, bold=True)
    D.bar_chart(s, 6.85, 2.2, 5.9, 4.3, [names[k] for k in wk], [wk[k] for k in wk], TEAL, "0%", "thresholds.confidence.weights")

    # 14 ---- appendix / methodology (dark close)
    s = D.slide(dark=True, notes="Methodology and limitations. All data is synthetic; no real organisation is represented.")
    D.title(s, "Appendix · methodology", "How to read these numbers, and their limits", dark=True)
    col = [
        [("Data", {"bold": True, "color": "7FC8BE", "size": 15})],
        f"{F['dq.records']} of {F['dq.total']} records usable; data quality {F['dq.score']}/100.",
        "Four injected defects were caught: a duplicate, a missing value, an invalid value and an inconsistent total. Each was handled by a documented exception; source data is never edited.",
        [("Trends", {"bold": True, "color": "7FC8BE", "size": 15})],
        f"At least {F['th.min_weeks']} weeks, Mann-Kendall |tau| ≥ {F['th.tau']} with p ≤ {F['th.p']}, a change of at least {F['th.min_pct']}% ({F['th.min_pp']} pp for rates), "
        f"and at least {F['th.effect']}× normal noise. Rates use pooled counts and a two-proportion test.",
    ]
    col2 = [
        [("Health index", {"bold": True, "color": "7FC8BE", "size": 15})],
        f"Relative: {F['health.base']} = organisation level in weeks 1–4, {F['health.better']} = {F['health.band']} better, {F['health.worse']} = {F['health.band']} worse. Not an absolute grade.",
        [("Causality", {"bold": True, "color": "7FC8BE", "size": 15})],
        f"Explanations are possibilities. They are rejected if a common driver explains them (partial r < {F['th.partial']}) or if the series only share a time trend.",
        [("Limitations", {"bold": True, "color": "7FC8BE", "size": 15})],
        "Synthetic data with one service per team and platform; 16 weeks is a short history; small weekly samples limit what rate changes can be proven.",
    ]
    D.text(s, X0, 1.95, 5.8, 4.8, col, 13, "D5D8DC", spacing=8)
    D.text(s, 6.95, 1.95, 5.8, 4.8, col2, 13, "D5D8DC", spacing=8)
    return D


# ======================================================================== report
def build_report(F, ctx):
    risks, recs, catalog = ctx["risks"], ctx["recs"], ctx["catalog"]
    L = [f"# Engineering Health Review",
         "",
         f"_Weeks 1–16 ({F['period.first']} to {F['period.last']}, week starts) · synthetic organisation · generated by Agent 10 from validated outputs._",
         "",
         "## Summary",
         "",
         f"Engineering health is **{F['kpi.overall.now']}/100**, {F['kpi.overall.delta']} points since week 1 (relative index; {F['health.base']} = the organisation's level in weeks 1–4). "
         f"{F['risk.count']} validated risks, {F['risk.critical']} critical.",
         "",
         f"_{STAT_NOTE}_",
         "",
         f"- **Atlas is slowing down and losing quality.** Cycle time {F['atlas.cycle.base']} → {F['atlas.cycle.recent']} ({F['atlas.cycle.chg']}), PR review time "
         f"{F['atlas.review.base']} → {F['atlas.review.recent']} ({F['atlas.review.chg']}), defects {F['atlas.defects.base']} → {F['atlas.defects.recent']} per week, "
         f"test coverage {F['atlas.coverage.base']} → {F['atlas.coverage.recent']}.",
         f"- **Data platform releases fail far more often.** Change failure rate {F['data.cfr.base']} → {F['data.cfr.recent']}, rollback rate {F['data.rollback.base']} → "
         f"{F['data.rollback.recent']}, incidents {F['data.incidents.base']} → {F['data.incidents.recent']} per week, on-call pages {F['data.oncall.base']} → {F['data.oncall.recent']} per week.",
         f"- **Phoenix improved.** Coverage {F['phoenix.coverage.base']} → {F['phoenix.coverage.recent']}, cycle time {F['phoenix.cycle.base']} → {F['phoenix.cycle.recent']}, "
         f"change failure rate {F['phoenix.cfr.base']} → {F['phoenix.cfr.recent']}. Its escaped-defect rate fell ({F['phoenix.esc.counts']}) but is not yet significant (p = {F['phoenix.esc.p']}).",
         f"- **Nova's week-11 spike was a one-off** (CI failure rate {F['nova.ci_failure_rate_pct.value']} vs an expected {F['nova.ci_failure_rate_pct.expected']}), back to normal the next week.",
         "",
         "## Health indices",
         "",
         "| Index | Now (4-week rolling) | Change since week 1 |",
         "|---|---|---|"]
    for k, label in (("overall", "Engineering health"), ("delivery", "Delivery"), ("quality", "Quality"), ("reliability", "Reliability"), ("efficiency", "Efficiency")):
        L.append(f"| {label} | {F[f'kpi.{k}.now']} | {F[f'kpi.{k}.delta']} |")
    L += ["", "## Risks", "", "| # | Severity | Risk | Category | Confidence | Validation |", "|---|---|---|---|---|---|"]
    for r in risks:
        L.append(f"| {r['rank']} | {r['severity']} | {r['risk']} | {r['category']} | {F[f'risk.{r['id']}.conf']} | {'Passed' if r['validation']['verdict'] == 'PASS' else 'Caveat'} |")
    L += ["", "Possible explanations that survived validation are possibilities, not causes. "
              f"For the Data platform, incidents attributed to changes rose {F['data.chg_rel.base']} → {F['data.chg_rel.recent']} (p = {F['data.chg_rel.p']}, a lead to check).",
          "", "## Recommended actions", "", "| Priority | Action | Owner | Success measure |", "|---|---|---|---|"]
    for r in recs:
        ms = r["measurement_of_success"]
        L.append(f"| {r['priority']} | {r['headline']} | {r['owner_type']} | {catalog[ms['metric']]['label']}: {F[f'rec.{r['id']}.target']} by {F[f'rec.{r['id']}.by']} |")
    L += ["", "## How these findings were produced", "",
          f"Ten agents with explicit contracts. Data quality {F['dq.score']}/100 ({F['dq.records']} of {F['dq.total']} records usable). "
          f"{F['verdict.total']} findings were re-checked by the evidence validator: {F['verdict.PASS']} passed, {F['verdict.WARN']} carry a caveat, {F['verdict.REJECT']} were rejected. "
          f"Example rejection: rising throughput at Titan correlates with longer builds (r = {F['rejected_example.r']}), but both follow headcount; controlling for it, r = {F['rejected_example.partial']}.",
          "",
          f"A trend needs at least {F['th.min_weeks']} weeks, a significant monotonic change (p ≤ {F['th.p']}), at least {F['th.min_pct']}% ({F['th.min_pp']} pp for rates) "
          f"and at least {F['th.effect']}× normal noise. All data is synthetic.", ""]
    return "\n".join(L)


# ======================================================================== checks
def deck_text(path):
    prs = Presentation(str(path))
    out = []
    for i, s in enumerate(prs.slides, start=1):
        for shp in s.shapes:
            if shp.has_text_frame:
                out.append((i, "shape", shp.text_frame.text))
            if shp.has_table:
                for row in shp.table.rows:
                    for cell in row.cells:
                        out.append((i, "table", cell.text))
        if s.has_notes_slide:
            out.append((i, "notes", s.notes_slide.notes_text_frame.text))
    return prs, out


def unsourced_numbers(texts, F):
    allowed = F.allowed_numbers()
    bad = []
    for where, kind, t in texts:
        for n in NUMBER.findall(t):
            x = float(n)
            if x.is_integer() and x <= 22:  # week numbers, slide numbers, counts of teams/agents
                continue
            if x not in allowed:
                bad.append({"where": where, "kind": kind, "number": n, "text": t[:90]})
    return bad


def consistency_checks(F, ctx, deck, report_text):
    checks = []

    def add(name, ok, detail):
        checks.append({"check": name, "result": "PASS" if ok else "FAIL", "detail": detail})

    prs, texts = deck_text(deck_path())
    add("deck_opens", len(prs.slides) == 14, f"{len(prs.slides)} slides")
    bad = unsourced_numbers(texts, F)
    add("deck_numbers_sourced", not bad, "every number on slides, tables and notes is in the fact sheet" if not bad else f"unsourced: {bad[:5]}")
    bad_r = unsourced_numbers([("report", "md", report_text)], F)
    add("report_numbers_sourced", not bad_r, "every number in the report is in the fact sheet" if not bad_r else f"unsourced: {bad_r[:5]}")

    # charts: series in the saved file equal their sources
    saved = []
    for i, s in enumerate(prs.slides, start=1):
        for shp in s.shapes:
            if shp.has_chart:
                for ser in shp.chart.plots[0].series:
                    saved.append((i, ser.name, list(ser.values)))
    mism = []
    for src in deck.chart_sources:
        got = next((v for (i, n, v) in saved if i == src["slide"] and n == src["series"]), None)
        exp = src["values"]
        if got is None or any((a is None) != (b is None) or (a is not None and abs(a - b) > 1e-9) for a, b in zip(got, exp)):
            mism.append(f"slide {src['slide']} {src['series']}")
    add("charts_match_sources", not mism and len(saved) == len(deck.chart_sources), f"{len(saved)} chart series identical to source" if not mism else f"mismatch: {mism}")

    # raw <-> analysis: recompute two headline facts straight from the CSVs
    validation = read_envelope(AGENTS["01"])
    data = load_clean_data(validation["analysis_directives"])
    dep = [r for r in data["deployments"] if r["platform"] == "Data" and r["week"] >= 13]
    raw_cfr = 100 * sum(r["status"] == "FAILED" for r in dep) / len(dep)
    add("raw_vs_analysis_cfr", abs(raw_cfr - F.facts["data.cfr.recent"].value) < 0.01,
        f"Data platform CFR weeks 13-16 from deployments.csv = {raw_cfr:.3f}% vs fact {F.facts['data.cfr.recent'].value}")
    defects = [sum(1 for r in data["defects"] if r["team"] == "Atlas" and r["week"] == w) for w in range(13, 17)]
    add("raw_vs_analysis_defects", median(defects) == F.facts["atlas.defects.recent"].value,
        f"Atlas defects/week weeks 13-16 median from defects.csv = {median(defects)} vs fact {F.facts['atlas.defects.recent'].value}")

    # qualitative claims in slide titles must hold in the data
    comp, cat = ctx["comp"], ctx["catalog"]
    teams = ["Atlas", "Nova", "Orion", "Phoenix", "Titan"]

    def worst(m):
        key = (lambda t: comp[(t, m)]) if cat[m]["polarity"] == "lower_is_better" else (lambda t: -comp[(t, m)])
        return max(teams, key=key)

    data_teams = {p.split("/")[0] for p in ctx["store"]["platform:Data"]["pairs"]}
    worst_cfr = sorted(teams, key=lambda t: -comp[(t, "change_failure_rate_pct")])[:2]
    claims = {
        "Atlas worst on cycle time, review time, coverage and defects": all(worst(m) == "Atlas" for m in
                                                                          ("cycle_time_days", "pr_review_time_hours", "test_coverage_pct", "defects")),
        "two worst change-failure teams are Data-platform teams": set(worst_cfr) <= data_teams,
        "timeline trends all last >= minimum trend length": all(f["evidence"]["computed"]["window_weeks"] >= ctx["th"]["trend"]["min_trend_weeks"]
                                                                for f in ctx["timeline"]),
        "Phoenix has no deteriorating trend": not any(f["type"] == "trend" and f["classification"] == "DETERIORATING"
                                                      and f["dimension"]["scope"] == "team:Phoenix" and f["id"] in ctx["ok"] for f in ctx["v08"].values()),
    }
    failed = [k for k, ok in claims.items() if not ok]
    add("title_claims_hold", not failed, f"{len(claims)} qualitative claims verified" if not failed else f"false claims: {failed}")

    # analysis <-> dashboard
    dash = agent09_dashboard.data_file()
    if dash.exists():
        t = dash.read_text(encoding="utf-8")
        d = json.loads(t[t.index("=") + 1:].strip().rstrip(";"))
        same = (d["kpi"]["org"]["overall"]["rolling"][-1] == F.facts["kpi.overall.now"].value
                and d["risks"] == [r["id"] for r in ctx["risks"]]
                and all(d["findings"][r["id"]]["confidence"]["score"] == r["confidence"]["score"] for r in ctx["risks"]))
        add("dashboard_agreement", same, "dashboard KPI, risk order and confidences identical to the deck" if same else "dashboard differs from deck - regenerate both")
    else:
        add("dashboard_agreement", False, "dashboard data missing - run Agent 09 first")
    return checks


def run():
    F, ctx = build_facts()
    deck = build_deck(F, ctx)
    deck_path().parent.mkdir(parents=True, exist_ok=True)
    deck.prs.core_properties.title = "Engineering Health Review"
    deck.prs.core_properties.author = "Engineering Health Intelligence System (Agent 10)"
    deck.prs.save(str(deck_path()))
    report = build_report(F, ctx)
    report_path().parent.mkdir(parents=True, exist_ok=True)
    report_path().write_text(report, encoding="utf-8", newline="\n")
    checks = consistency_checks(F, ctx, deck, report)
    status = "PASS" if all(c["result"] == "PASS" for c in checks) else "FAIL"
    summary = {"slides": len(deck.prs.slides), "facts": len(F.facts), "chart_series": len(deck.chart_sources),
               "deck": deck_path().name, "report": report_path().name}
    deps = [AGENTS[k] for k in ("01", "02", "03", "04", "05", "06", "07", "08")]
    write_envelope(AGENT, status, deps, [envelope_path(a) for a in deps] + [processed_dir() / "metric_series.json"],
                   summary, [], consistency_checks=checks, fact_sheet=F.manifest(), chart_sources=deck.chart_sources,
                   errors=[c["detail"] for c in checks if c["result"] == "FAIL"])
    return summary, checks, status


def main():
    s, checks, status = run()
    print(f"[{AGENT}] status={status} slides={s['slides']} facts={s['facts']} chart_series={s['chart_series']} -> {s['deck']}, {s['report']}")
    for c in checks:
        print(f"  [{c['result']}] {c['check']}: {c['detail'][:160]}")


if __name__ == "__main__":
    main()
