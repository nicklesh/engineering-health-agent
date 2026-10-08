"""Phase 6 tests: Agent 10 (executive report + PowerPoint).

The brief's output-consistency requirement: raw data <-> analysis <-> dashboard <-> PowerPoint.
"""
import shutil
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

from src.common import envelope
from src.config import ROOT
from src.reporting import agent09_dashboard, agent10_executive_report as a10
from src.reporting.facts import FactSheet
from tests.pipeline_helper import shared_run


class ReportCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp(prefix="ehis-report-"))
        (cls.tmp / "dashboard").mkdir()
        shutil.copy(ROOT / "dashboard" / "index.html", cls.tmp / "dashboard" / "index.html")
        cls._dash, cls._out = agent09_dashboard.DASHBOARD_DIR, a10.OUTPUT_DIR
        agent09_dashboard.DASHBOARD_DIR = cls.tmp / "dashboard"
        a10.OUTPUT_DIR = cls.tmp / "output"
        envelope.set_processed_dir(shared_run())
        try:
            agent09_dashboard.run()
            cls.summary, cls.checks, cls.status = a10.run()
            cls.F, cls.ctx = a10.build_facts()
        finally:
            envelope.set_processed_dir(None)
            agent09_dashboard.DASHBOARD_DIR, a10.OUTPUT_DIR = cls._dash, cls._out

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)


class TestAgent10(ReportCase):
    def test_all_consistency_checks_pass(self):
        self.assertEqual(self.status, "PASS", [c for c in self.checks if c["result"] != "PASS"])
        names = {c["check"] for c in self.checks}
        for required in ("deck_opens", "deck_numbers_sourced", "report_numbers_sourced", "charts_match_sources",
                         "raw_vs_analysis_cfr", "raw_vs_analysis_defects", "dashboard_agreement", "title_claims_hold"):
            self.assertIn(required, names)

    def test_deck_structure(self):
        prs = Presentation(str(self.tmp / "output" / "presentation" / "engineering_health_review.pptx"))
        self.assertTrue(10 <= len(prs.slides) <= 14)
        for i, s in enumerate(prs.slides, start=1):
            with self.subTest(slide=i):
                self.assertTrue(s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip(), "speaker notes present")
        charts = sum(1 for s in prs.slides for shp in s.shapes if shp.has_chart)
        self.assertGreaterEqual(charts, 5, "native, editable charts")

    def test_rejected_findings_never_reach_the_deck(self):
        prs, texts = a10.deck_text(self.tmp / "output" / "presentation" / "engineering_health_review.pptx")
        all_text = " ".join(t for _, _, t in texts)
        # The Titan correlation appears only as the labelled REJECTED example, never as a finding or explanation.
        self.assertIn("REJECTED EXAMPLE", all_text)
        self.assertNotIn("may be contributing to the change in build time", all_text)
        self.assertNotIn("on-call pages may be contributing", all_text)

    def test_report_written(self):
        text = (self.tmp / "output" / "reports" / "executive_report.md").read_text(encoding="utf-8")
        self.assertIn("## Risks", text)
        self.assertIn(self.F["kpi.overall.now"], text)

    def test_unsourced_number_is_caught(self):
        F = FactSheet()
        F.add("x", 27.1, "27.1%", "test")
        bad = a10.unsourced_numbers([(1, "shape", "Failure rate rose to 27.1%, costing 3.7 million")], F)
        self.assertEqual([b["number"] for b in bad], ["3.7"])
        self.assertEqual(a10.unsourced_numbers([(1, "shape", "Week 16, slide 3: 27.1%")], F), [])

    def test_facts_trace_to_validated_findings(self):
        """Every fact comes from a validated finding - except those explicitly marked as the
        labelled rejection example, which must come from a REJECTED finding."""
        v08 = self.ctx["v08"]
        for f in self.F.manifest():
            src = f["source"].split(".")[0]
            if src.startswith(("TRD-", "ANO-", "RISK-", "REC-")):
                with self.subTest(fact=f["key"]):
                    self.assertIn(src, v08)
                    rejected = v08[src]["validation"]["verdict"] == "REJECT"
                    self.assertEqual(rejected, f["key"].startswith("rejected_example."))

    def test_rejected_example_only_on_evidence_slide(self):
        prs, texts = a10.deck_text(self.tmp / "output" / "presentation" / "engineering_health_review.pptx")
        partial = self.F["rejected_example.partial"]
        slides = {i for i, kind, t in texts if partial in t and kind != "notes"}
        evidence_slide = next(i for i, kind, t in texts if "Evidence & confidence model".upper() in t)
        self.assertEqual(slides, {evidence_slide})


if __name__ == "__main__":
    unittest.main()
